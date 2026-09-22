import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from app import web
from app.catalog_db import CatalogDatabase
from app.services.excel_product_catalog import ExcelProductCatalog


class ProductsReportTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.database_path = Path(self.temp.name) / "catalog.db"
        database = CatalogDatabase(self.database_path)
        database.initialize()
        with database.transaction() as connection:
            connection.execute(
                "INSERT INTO catalog_excel_batches (id,file_sha256,source_filename,row_count,"
                "total_stock,positive_rows,zero_rows,status,created_at,applied_at) "
                "VALUES ('report','sha','report.xlsx',0,0,0,0,'active',?,?)",
                ("2026-09-01T09:00:00+00:00", "2026-09-01T09:00:00+00:00"),
            )
        catalog = ExcelProductCatalog(database)
        self.matches = []
        for index in range(123):
            self.matches.append(catalog.create_product(
                name="Report Watch {:03d}".format(index),
                model="Report Model", article="REP-{:03d}".format(index),
                brand="Report Brand", category="Часы", cell="A-01",
                stock=2 if index else 0, price="1990.50",
            ))
        self.unrelated = catalog.create_product(
            name="Other Product", model="Other", article="OTHER",
            brand="Other Brand", category="Аксессуары", cell="B-02", stock=4,
        )
        self.inactive = catalog.create_product(
            name="Archived Report Watch", model="Report Model", article="OLD-1",
            brand="Report Brand", category="Часы", cell="A-01", stock=1,
        )
        with database.transaction() as connection:
            connection.execute(
                "UPDATE catalog_excel_products SET active=0 WHERE id=?",
                (self.inactive["id"],),
            )
        self.environment = mock.patch.dict(
            "os.environ", {"CATALOG_DATABASE_PATH": str(self.database_path)}
        )
        self.environment.start()
        web.app.config.update(TESTING=True, AUTH_TESTING=False)
        self.client = web.app.test_client()

    def tearDown(self):
        self.environment.stop()
        self.temp.cleanup()

    def test_products_page_has_report_button_immediately_before_add(self):
        response = self.client.get("/app/products")
        self.assertEqual(response.status_code, 200)
        markup = response.get_data(as_text=True)
        report = markup.index('href="/app/products/report"')
        add = markup.index('id="openWarehouseAddModal"')
        self.assertLess(report, add)
        self.assertNotIn("Карта склада", markup[report:add])
        self.assertNotIn("PDF", markup[report:add])

    def test_report_filters_kpis_status_and_responsive_controls(self):
        active = self.client.get(
            "/app/products/report?q=Report&stock_state=in&activity=active&per_page=50"
        )
        self.assertEqual(active.status_code, 200)
        markup = active.get_data(as_text=True)
        self.assertIn("Отчёт по товарам", markup)
        self.assertIn('<strong id="report-total">122</strong>', markup)
        self.assertIn("Найдено товаров: 122", markup)
        self.assertIn('<strong id="report-in-stock">122</strong>', markup)
        self.assertIn('<strong id="report-out-of-stock">0</strong>', markup)
        self.assertIn('<strong id="report-stock-total">244</strong>', markup)
        self.assertIn("data-erp-table-key=\"products-report\"", markup)
        self.assertIn("data-report-column=\"photo\"", markup)
        self.assertIn("@media (max-width: 767px)", markup)
        self.assertEqual(markup.count("Report Watch "), 50)

        all_statuses = self.client.get("/app/products/report?activity=all")
        self.assertIn(
            '<strong id="report-total">125</strong>',
            all_statuses.get_data(as_text=True),
        )

        out_of_stock = self.client.get(
            "/app/products/report?q=Report&stock_state=out&activity=active"
        )
        self.assertIn(
            '<strong id="report-total">1</strong>',
            out_of_stock.get_data(as_text=True),
        )

        inactive = self.client.get("/app/products/report?q=Archived&activity=inactive")
        inactive_markup = inactive.get_data(as_text=True)
        self.assertIn('<strong id="report-total">1</strong>', inactive_markup)
        self.assertIn("Неактивен", inactive_markup)

        product = self.matches[1]
        brand_only = self.client.get(
            "/app/products/report?brand_id={}".format(product["brand_id"])
        )
        self.assertIn(
            '<strong id="report-total">123</strong>',
            brand_only.get_data(as_text=True),
        )
        combined = self.client.get(
            "/app/products/report?brand_id={}&category_id={}&model_id={}&cell=A-01"
            "&stock_state=in&activity=active".format(
                product["brand_id"], product["category_id"], product["model_id"]
            )
        )
        self.assertIn('<strong id="report-total">122</strong>', combined.get_data(as_text=True))

    def test_filtered_count_equals_excel_rows_independent_of_pagination(self):
        query = "q=Report&stock_state=in&activity=active&per_page=50&page=2"
        page = self.client.get("/app/products/report?" + query)
        self.assertEqual(page.status_code, 200)
        self.assertIn('<strong id="report-total">122</strong>', page.get_data(as_text=True))

        export = self.client.get("/app/products/report.xlsx?" + query)
        self.assertEqual(export.status_code, 200)
        workbook = load_workbook(BytesIO(export.data), read_only=True)
        sheet = workbook.active
        self.assertEqual(sheet.max_row - 1, 122)
        self.assertEqual(sheet["A1"].value, "Фото")
        self.assertEqual(sheet["K1"].value, "Ячейка")
        self.assertEqual(sheet["C2"].value, 2)
        self.assertEqual(sheet["D2"].value, 1990.5)

        all_export = self.client.get("/app/products/report.xlsx")
        all_sheet = load_workbook(BytesIO(all_export.data), read_only=True).active
        self.assertEqual(all_sheet.max_row - 1, 124)

        inactive_export = self.client.get(
            "/app/products/report.xlsx?activity=inactive"
        )
        inactive_sheet = load_workbook(
            BytesIO(inactive_export.data), read_only=True
        ).active
        self.assertEqual(inactive_sheet.max_row - 1, 1)


if __name__ == "__main__":
    unittest.main()
