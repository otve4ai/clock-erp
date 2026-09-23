import tempfile
import unittest
import inspect
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from datetime import date
from io import BytesIO
from pathlib import Path

from app.catalog_db import CatalogDatabase
from app.schema_migrations import apply_migrations
from app.services.excel_product_catalog import ExcelProductBatchService, ExcelProductCatalog
from app.services.business_analytics import BusinessAnalytics, parse_filters
from app.services.product_excel_export import ProductExcelExport
from app.services.product_bundles import ProductBundles
from app.services.warehouse_stock import WarehouseStockError, WarehouseStockService, set_balance


def row(number, article, stock):
    return {"excel_row": number, "excel_name": article, "excel_name_raw": article,
            "excel_article": article, "excel_brand": "Test", "category": "Часы",
            "stock": stock, "stock_valid": True, "cell": "A-1",
            "match_status": "not_found", "match_method": "none",
            "confidence": 0, "alternatives": []}


class CountingDatabase:
    def __init__(self, database):
        self.database = database
        self.queries = 0

    def initialize(self):
        return self.database.initialize()

    @contextmanager
    def connect(self):
        with self.database.connect() as connection:
            owner = self
            class Proxy:
                def execute(self, *args, **kwargs):
                    owner.queries += 1
                    return connection.execute(*args, **kwargs)
            yield Proxy()


class MultiwarehouseStage5Test(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / "catalog.db"
        apply_migrations(self.path, app_commit="stage5-test")
        self.db = CatalogDatabase(self.path)
        ExcelProductBatchService(self.db).apply([
            row(2, "HEAD", 10), row(3, "STRAP", 4),
            row(4, "BUNDLE", 0), row(5, "ORDINARY", 5),
        ], "b" * 64, "stage5.xlsx")
        self.service = WarehouseStockService(self.db)
        warehouses = {item["code"]: item for item in self.service.list_warehouses()}
        self.udelnaya = warehouses["udelnaya"]["id"]
        self.hong_kong = warehouses["hong-kong"]["id"]
        with self.db.connect() as connection:
            self.products = {item["excel_article"]: item["id"] for item in connection.execute(
                "SELECT id,excel_article FROM catalog_excel_products"
            )}

    def tearDown(self):
        self.temp.cleanup()

    def stock(self, article, warehouse, value):
        with self.db.transaction() as connection:
            set_balance(connection, self.products[article], warehouse, value)

    def test_default_selection_is_only_udelnaya(self):
        self.assertEqual(self.service.selected_warehouse_ids("new-user"), [self.udelnaya])

    def test_warehouse_api_model_has_required_fields_and_totals(self):
        rows = self.service.list_for_user("new-user")
        self.assertEqual({"id", "code", "name", "selected", "total"} - set(rows[0]), set())
        self.assertTrue(next(item for item in rows if item["id"] == self.udelnaya)["selected"])

    def test_new_warehouse_is_not_automatically_selected(self):
        created = self.service.create_warehouse("Новый", "new")
        self.assertNotIn(created["id"], self.service.selected_warehouse_ids("new-user"))

    def test_preferences_are_isolated_and_persistent(self):
        self.service.set_selected_warehouses("a", [self.hong_kong])
        self.assertEqual(self.service.selected_warehouse_ids("a"), [self.hong_kong])
        self.assertEqual(self.service.selected_warehouse_ids("b"), [self.udelnaya])

    def test_empty_selection_is_rejected(self):
        with self.assertRaises(WarehouseStockError):
            self.service.set_selected_warehouses("a", [])

    def test_rename_and_duplicate_name_validation(self):
        renamed = self.service.rename_warehouse(self.hong_kong, "Гонконг 2")
        self.assertEqual(renamed["name"], "Гонконг 2")
        with self.assertRaises(WarehouseStockError):
            self.service.rename_warehouse(self.hong_kong, "Удельная")

    def test_archive_only_empty_non_default_warehouse(self):
        reserve = self.service.create_warehouse("Резерв", "reserve")
        self.service.archive_warehouse(reserve["id"])
        self.assertNotIn(reserve["id"], [item["id"] for item in self.service.list_warehouses()])
        with self.assertRaises(WarehouseStockError):
            self.service.archive_warehouse(self.udelnaya)

    def test_nonempty_warehouse_cannot_be_archived(self):
        self.stock("ORDINARY", self.hong_kong, 1)
        with self.assertRaises(WarehouseStockError):
            self.service.archive_warehouse(self.hong_kong)

    def test_ordinary_total_uses_selected_warehouses(self):
        self.stock("ORDINARY", self.hong_kong, 3)
        product = ExcelProductCatalog(self.db).list_products(
            product_id=self.products["ORDINARY"], warehouse_ids=[self.udelnaya, self.hong_kong],
            include_facets=False,
        )["items"][0]
        self.assertEqual(product["stock"], 8)
        self.assertEqual([item["quantity"] for item in product["warehouse_breakdown"]], [5, 3])

    def test_bundle_is_calculated_per_warehouse_then_summed(self):
        ProductBundles(self.db).configure(self.products["BUNDLE"], [
            {"component_id": self.products["HEAD"], "quantity": 1},
            {"component_id": self.products["STRAP"], "quantity": 1},
        ])
        self.stock("HEAD", self.hong_kong, 2)
        self.stock("STRAP", self.hong_kong, 8)
        product = ExcelProductCatalog(self.db).list_products(
            product_id=self.products["BUNDLE"], warehouse_ids=[self.udelnaya, self.hong_kong],
            include_facets=False,
        )["items"][0]
        self.assertEqual(product["stock"], 6)
        self.assertEqual([item["quantity"] for item in product["warehouse_breakdown"]], [4, 2])

    def test_filter_sort_and_pagination_use_display_stock(self):
        self.stock("ORDINARY", self.hong_kong, 20)
        catalog = ExcelProductCatalog(self.db).list_products(
            warehouse_ids=[self.hong_kong], stock_state="in", sort_by="stock",
            sort_dir="desc", page=1, per_page=1, include_facets=False,
        )
        self.assertEqual(catalog["total"], 1)
        self.assertEqual(catalog["items"][0]["excel_article"], "ORDINARY")

    def test_out_of_stock_filter_uses_selected_warehouse(self):
        catalog = ExcelProductCatalog(self.db).list_products(
            warehouse_ids=[self.hong_kong], stock_state="out", include_facets=False,
        )
        self.assertEqual({item["excel_article"] for item in catalog["items"]},
                         {"HEAD", "STRAP", "BUNDLE", "ORDINARY"})

    def test_pagination_is_applied_after_warehouse_filter(self):
        self.stock("HEAD", self.hong_kong, 1)
        self.stock("STRAP", self.hong_kong, 1)
        first = ExcelProductCatalog(self.db).list_products(
            warehouse_ids=[self.hong_kong], stock_state="in", sort_by="name",
            page=1, per_page=1, include_facets=False,
        )
        second = ExcelProductCatalog(self.db).list_products(
            warehouse_ids=[self.hong_kong], stock_state="in", sort_by="name",
            page=2, per_page=1, include_facets=False,
        )
        self.assertEqual(first["total"], 2)
        self.assertNotEqual(first["items"][0]["id"], second["items"][0]["id"])

    def test_product_read_model_has_no_legacy_stock_predicates(self):
        source = inspect.getsource(ExcelProductCatalog.list_products)
        self.assertNotIn('where.append("p.stock', source)
        self.assertNotIn('"stock": "p.stock"', source)

    def test_top_units_do_not_double_count_virtual_bundle(self):
        ProductBundles(self.db).configure(self.products["BUNDLE"], [
            {"component_id": self.products["HEAD"], "quantity": 1},
            {"component_id": self.products["STRAP"], "quantity": 1},
        ])
        counts = ExcelProductCatalog(self.db).stock_tab_counts([self.udelnaya])
        self.assertEqual(counts["units_total"], 19)
        self.assertEqual(counts["in_stock"], 4)

    def test_page_query_count_does_not_grow_with_rows(self):
        measured = CountingDatabase(self.db)
        ExcelProductCatalog(measured).list_products(
            warehouse_ids=[self.udelnaya, self.hong_kong], per_page=100,
            include_facets=False,
        )
        self.assertLessEqual(measured.queries, 6)

    def test_frontend_contracts_cover_selector_stock_popup_and_transfer(self):
        root = Path(__file__).resolve().parents[1]
        template = (root / "app/templates/warehouse.html").read_text(encoding="utf-8")
        script = (root / "app/static/js/warehouse-multiwarehouse.js").read_text(encoding="utf-8")
        self.assertIn("🏭 Склады:", template)
        self.assertIn("warehouse-stock-popover", template)
        self.assertIn("warehouseTransferDialog", template)
        self.assertIn("/api/v1/warehouse-preferences", script)
        self.assertIn("/api/v1/warehouse-transfers", script)

    def test_operational_forms_have_explicit_warehouse_selector(self):
        root = Path(__file__).resolve().parents[1]
        sales = (root / "app/templates/sales.html").read_text(encoding="utf-8")
        writeoffs = (root / "app/templates/_writeoffs.html").read_text(encoding="utf-8")
        inventory = (root / "app/static/js/warehouse-operation-selectors.js").read_text(encoding="utf-8")
        receipts = (root / "app/templates/excel_receipt_preview.html").read_text(encoding="utf-8")
        bundle = (root / "app/templates/product_bundle.html").read_text(encoding="utf-8")
        self.assertIn('name="warehouse_id"', sales)
        self.assertIn('name="warehouse_id"', writeoffs)
        self.assertIn("startForm", inventory)
        self.assertIn('name="warehouse_id"', receipts)
        self.assertIn('name="warehouse_id"', bundle)
        self.assertNotIn("Остаток Удельной", bundle)

    def test_analytics_stock_uses_selected_warehouses(self):
        self.stock("ORDINARY", self.hong_kong, 7)
        filters = parse_filters({"stock_state": "positive"}, today=date(2026, 9, 23))
        with self.db.connect() as connection:
            udelnaya = BusinessAnalytics(
                self.db, warehouse_ids=[self.udelnaya]
            )._stock_rows(connection, filters)["stock_all_rows"]
            hong_kong = BusinessAnalytics(
                self.db, warehouse_ids=[self.hong_kong]
            )._stock_rows(connection, filters)["stock_all_rows"]
        self.assertEqual(next(row for row in udelnaya if row["id"] == self.products["ORDINARY"])["stock"], 5)
        self.assertEqual(next(row for row in hong_kong if row["id"] == self.products["ORDINARY"])["stock"], 7)

    def test_export_contains_dynamic_warehouse_breakdown(self):
        self.stock("ORDINARY", self.hong_kong, 3)
        item = ExcelProductCatalog(self.db).list_products(
            product_id=self.products["ORDINARY"],
            warehouse_ids=[self.udelnaya, self.hong_kong], include_facets=False,
        )["items"][0]
        payload = ProductExcelExport(self.db).build(
            [item], 1, warehouses=[
                {"id": self.udelnaya, "name": "Удельная"},
                {"id": self.hong_kong, "name": "Гонконг"},
            ],
        )
        from openpyxl import load_workbook
        sheet = load_workbook(BytesIO(payload), read_only=True).active
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows[0][8:11], ("Остаток", "Удельная", "Гонконг"))
        self.assertEqual(rows[1][8:11], (8, 5, 3))

    def test_concurrent_transfers_cannot_overspend_source(self):
        def transfer(key):
            try:
                self.service.transfer(
                    self.products["ORDINARY"], self.udelnaya, self.hong_kong, 3,
                    idempotency_key=key,
                )
                return True
            except WarehouseStockError:
                return False
        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(executor.map(transfer, ("concurrent-a", "concurrent-b")))
        self.assertEqual(results.count(True), 1)
        self.assertEqual(self.service.get_balance(self.products["ORDINARY"], self.udelnaya), 2)
        self.assertEqual(self.service.get_balance(self.products["ORDINARY"], self.hong_kong), 3)

    def test_admin_actions_and_transfer_are_audited(self):
        actor = {"actor_id": "7", "actor_name": "Максим"}
        reserve = self.service.create_warehouse("Резерв QA", "reserve-qa", actor=actor)
        self.service.rename_warehouse(reserve["id"], "Резерв QA 2", actor=actor)
        self.service.archive_warehouse(reserve["id"], actor=actor)
        self.service.transfer(
            self.products["ORDINARY"], self.udelnaya, self.hong_kong, 1,
            actor=actor, idempotency_key="audit-transfer",
        )
        with self.db.connect() as connection:
            actions = [row[0] for row in connection.execute(
                "SELECT action FROM erp_audit_events WHERE actor_id='7' ORDER BY id"
            ).fetchall()]
        self.assertEqual(actions, ["created", "updated", "archived", "updated"])

    def test_database_constraints_and_indexes_cover_stock_keys(self):
        with self.db.connect() as connection:
            foreign_keys = {
                (row["from"], row["table"])
                for row in connection.execute(
                    "PRAGMA foreign_key_list(erp_product_warehouse_stock)"
                )
            }
            indexes = {
                row["name"] for row in connection.execute(
                    "PRAGMA index_list(erp_product_warehouse_stock)"
                )
            }
        self.assertIn(("product_id", "catalog_excel_products"), foreign_keys)
        self.assertIn(("warehouse_id", "erp_warehouses"), foreign_keys)
        self.assertTrue(any("warehouse" in name for name in indexes))


if __name__ == "__main__":
    unittest.main()
