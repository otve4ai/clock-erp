import tempfile
import unittest
from datetime import datetime
from io import BytesIO
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from app import web
from app.catalog_db import CatalogDatabase
from app.services.bitrix_catalog_importer import BitrixCatalogImporter
from app.services.excel_product_catalog import ExcelProductCatalog
from app.services import product_excel_export
from app.services.product_excel_export import ProductExcelExport


class ProductExcelExportTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.database_path = Path(self.temp.name) / "catalog.db"
        database = CatalogDatabase(self.database_path)
        database.initialize()
        with database.transaction() as connection:
            connection.execute(
                "INSERT INTO catalog_excel_batches (id,file_sha256,source_filename,row_count,"
                "total_stock,positive_rows,zero_rows,status,created_at,applied_at) "
                "VALUES ('export','sha','export.xlsx',0,0,0,0,'active',?,?)",
                ("2026-08-18T09:00:00+00:00", "2026-08-18T09:00:00+00:00"),
            )
        catalog = ExcelProductCatalog(database)
        self.ziiiro = catalog.create_product(
            name="=Опасная формула", model="Celeste", article="@SKU",
            brand="Ziiiro", category="Часы", stock=3, price="1234.50",
        )
        self.other = catalog.create_product(
            name="Обычные часы", model="Classic", article="SAFE",
            brand="Other", category="Часы", stock=3,
        )
        BitrixCatalogImporter(database).import_products([{
            "external_source": "bitrix",
            "external_product_id": "warehouse-product",
            "external_sku": "@SKU",
            "name": "=Опасная формула",
            "brand": "Ziiro",
            "category": "Часы",
            "active": True,
            "warehouse_stocks": [
                {"name": "Удельная", "quantity": 2},
                {"name": "Москва", "quantity": 1},
            ],
        }], "full_sync")
        with database.transaction() as connection:
            linked_id = connection.execute(
                "SELECT id FROM catalog_products WHERE external_product_id=?",
                ("warehouse-product",),
            ).fetchone()[0]
            connection.execute(
                "UPDATE catalog_excel_products SET bitrix_catalog_product_id=?, "
                "bitrix_description=?,updated_at=? WHERE id=?",
                (
                    linked_id, "+Описание",
                    "2026-08-18T10:30:00+00:00", self.ziiiro["id"],
                ),
            )
        self.environment = mock.patch.dict(
            "os.environ", {"CATALOG_DATABASE_PATH": str(self.database_path)}
        )
        self.environment.start()
        web.app.config.update(TESTING=True, AUTH_TESTING=False)
        self.client = web.app.test_client()

    def tearDown(self):
        product_excel_export._cached_available_warehouses.cache_clear()
        self.environment.stop()
        self.temp.cleanup()

    def test_available_warehouses_are_cached_until_catalog_changes(self):
        exporter = ProductExcelExport(CatalogDatabase(self.database_path))
        product_excel_export._cached_available_warehouses.cache_clear()
        with mock.patch.object(
            product_excel_export,
            "_read_available_warehouses",
            wraps=product_excel_export._read_available_warehouses,
        ) as reader:
            self.assertEqual(exporter.available_warehouses(), ["Москва", "Удельная"])
            self.assertEqual(exporter.available_warehouses(), ["Москва", "Удельная"])
            self.assertEqual(reader.call_count, 1)

            with exporter.database.transaction() as connection:
                connection.execute(
                    "UPDATE catalog_products SET normalized_payload_json = ? "
                    "WHERE external_product_id = ?",
                    (
                        '{"warehouse_stocks":[{"name":"Новый склад","quantity":1}]}',
                        "warehouse-product",
                    ),
                )

            self.assertEqual(exporter.available_warehouses(), ["Новый склад"])
            self.assertEqual(reader.call_count, 2)

    def workbook(self, url, data=None):
        response = (
            self.client.post(url, data=data)
            if data is not None else self.client.get(url)
        )
        self.assertEqual(
            response.status_code, 200,
            response.get_data(as_text=True) if response.status_code != 200 else "",
        )
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment; filename=", response.headers["Content-Disposition"])
        return response, load_workbook(BytesIO(response.data))

    def test_filtered_and_all_exports_ignore_pagination_and_keep_excel_types(self):
        response, workbook = self.workbook(
            "/app/products/export.xlsx?scope=filtered&q=Celeste&per_page=1"
            "&fields=sku&fields=name&fields=price&fields=stock_total"
            "&fields=updated_at"
        )
        sheet = workbook.active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:E2")
        values = [cell.value for cell in sheet[2]]
        self.assertEqual(values[:2], ["'@SKU", "'=Опасная формула"])
        self.assertEqual(values[2:4], [1234.5, 3])
        self.assertIsInstance(values[4], datetime)
        self.assertEqual(response.headers["X-Export-Count"], "1")

        _, all_workbook = self.workbook(
            "/app/products/export.xlsx?scope=all&q=Celeste&per_page=1"
        )
        self.assertEqual(all_workbook.active.max_row, 3)
        self.assertEqual(all_workbook.active.auto_filter.ref, "A1:I3")

    def test_dynamic_warehouse_columns_and_formula_injection_are_safe(self):
        _, workbook = self.workbook(
            "/app/products/export.xlsx?scope=filtered&q=Celeste"
            "&fields=sku&fields=name&fields=description&fields=stock_total"
            "&fields=warehouse:Москва&fields=warehouse:Удельная"
        )
        sheet = workbook.active
        self.assertEqual(
            [cell.value for cell in sheet[1]],
            ["Артикул / SKU", "Название", "Описание", "Общий остаток", "Москва", "Удельная"],
        )
        self.assertEqual(
            [cell.value for cell in sheet[2]],
            ["'@SKU", "'=Опасная формула", "'+Описание", 3, 1, 2],
        )
        self.assertTrue(all(cell.data_type != "f" for cell in sheet[2]))

    def test_site_issue_filter_is_applied_to_export_and_described_in_ui(self):
        with CatalogDatabase(self.database_path).transaction() as connection:
            connection.execute(
                "UPDATE catalog_excel_products "
                "SET bitrix_external_product_id=?, bitrix_active=0 WHERE id=?",
                ("warehouse-product", self.ziiiro["id"]),
            )

        _, workbook = self.workbook(
            "/app/products/export.xlsx?scope=filtered"
            "&site_issue=in_stock_inactive&fields=name"
        )
        self.assertEqual(workbook.active.max_row, 2)
        self.assertEqual(workbook.active["A2"].value, "'=Опасная формула")

        page = self.client.get(
            "/app/products?site_issue=in_stock_inactive"
        )
        markup = page.get_data(as_text=True)
        self.assertIn("Статус сайта: с остатком выключены", markup)

        javascript = (
            Path(web.app.root_path) / "static" / "js" / "product-export.js"
        ).read_text(encoding="utf-8")
        self.assertIn('currentState?.site_issue || params.get("site_issue")', javascript)
        self.assertIn("Статус сайта: С остатком выключены", javascript)

    def test_selected_export_mode_is_removed(self):
        response = self.client.post("/app/products/export.xlsx", data={
            "scope": "selected",
            "selected_ids": str(self.other["id"]),
            "fields": ["sku", "name"],
            "filename": "Товары_выбранные.xlsx",
        })
        self.assertEqual(response.status_code, 400)
        self.assertIn("режим", response.get_json()["message"])

    def test_post_rejects_removed_scope_and_requires_at_least_one_field(self):
        removed_scope = self.client.post(
            "/app/products/export.xlsx", data={"scope": "selected", "fields": "name"}
        )
        self.assertEqual(removed_scope.status_code, 400)
        self.assertIn("режим", removed_scope.get_json()["message"])
        no_fields = self.client.post(
            "/app/products/export.xlsx", data={"scope": "all"}
        )
        self.assertEqual(no_fields.status_code, 400)
        self.assertIn("поле", no_fields.get_json()["message"])

    def test_unauthenticated_user_cannot_download_catalogue(self):
        web.app.config["AUTH_TESTING"] = True
        try:
            response = self.client.get("/app/products/export.xlsx?scope=all")
        finally:
            web.app.config["AUTH_TESTING"] = False
        self.assertEqual(response.status_code, 302)
        self.assertIn("/login", response.headers["Location"])

    def test_export_modal_keeps_all_and_filtered_without_table_selection(self):
        page = self.client.get(
            "/app/products?stock_state=in&q=Celeste"
        )
        markup = page.get_data(as_text=True)
        self.assertIn("Экспорт в Excel", markup)
        self.assertIn("Экспорт товаров в Excel", markup)
        self.assertIn("Сейчас применены фильтры", markup)
        self.assertIn("Поиск: Celeste", markup)
        self.assertIn("Наличие: В наличии", markup)
        self.assertIn('name="scope" value="all"', markup)
        self.assertIn('name="scope" value="filtered"', markup)
        self.assertNotIn('name="scope" value="selected"', markup)
        self.assertNotIn("data-export-select", markup)
        self.assertNotIn("Только выбранные", markup)
        self.assertIn("warehouse:Москва", markup)

    def test_export_enrichment_chunks_more_than_sqlite_variable_limit(self):
        products = [{"id": product_id, "stock": 1} for product_id in range(1, 1002)]
        enriched = ProductExcelExport(
            CatalogDatabase(self.database_path)
        ).enrich(products)
        self.assertEqual(len(enriched), 1001)
        self.assertTrue(all("_export_warehouses" in item for item in enriched))


if __name__ == "__main__":
    unittest.main()
