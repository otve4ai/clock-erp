"""Standalone UI and local-only controls for CDEK payout reconciliation."""
import io
import zipfile
from datetime import datetime, timedelta
from decimal import Decimal

from flask import abort, jsonify, redirect, render_template, request, url_for

from app.clients.cdek import CdekError
from app.services.cdek_payouts import LABELS, TRACK, amount, iso_day, money
from app.services.cdek_delivery import display_time
from app.services.cdek_sync import CdekSync


def import_report(payload):
    """Read only shipment identifiers and dates, not private recipient columns."""
    import openpyxl
    if not payload or len(payload) > 10 * 1024 * 1024:
        raise ValueError("Допустим XLSX до 10 МБ.")
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as z:
            if sum(i.file_size for i in z.infolist()) > 40 * 1024 * 1024:
                raise ValueError("Слишком большой распакованный отчёт.")
        book = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        try:
            sheet = book.active
            if sheet.max_column > 100 or sheet.max_row > 10001:
                raise ValueError("Допустим отчёт до 10 000 строк и 100 столбцов.")
            rows = iter(sheet.values)
            header = list(next(rows))
            track_index = header.index("Номер заказа")
            date_index = header.index("Дата накладной")
            tracks, dates = set(), []
            for row in rows:
                value = row[track_index]
                if value is None:
                    continue
                track = str(int(value)) if isinstance(value, (int, float)) and value == int(value) else str(value).strip()
                if not TRACK.fullmatch(track):
                    raise ValueError("В отчёте есть некорректный номер накладной.")
                day = row[date_index]
                if isinstance(day, datetime):
                    day = day.date().isoformat()
                else:
                    day = datetime.strptime(str(day), "%d.%m.%Y").date().isoformat()
                tracks.add(track)
                dates.append(day)
            if not tracks:
                raise ValueError("В отчёте нет накладных.")
            return sorted(tracks), min(dates)
        finally:
            book.close()
    except (OSError, KeyError, StopIteration, zipfile.BadZipFile):
        raise ValueError("Не удалось прочитать отчёт заказов СДЭК в формате XLSX.") from None


def build_view(snapshot, args, today):
    start = args.get("from") or (today-timedelta(days=44)).isoformat()
    end = args.get("to") or today.isoformat()
    try:
        start, end = iso_day(start), iso_day(end)
        if start > end or end > today.isoformat():
            raise ValueError()
    except (ValueError, CdekError):
        raise ValueError("Проверьте даты: начало не позже конца, конец не позже сегодняшнего дня.") from None
    kind = args.get("date_kind", "created")
    if kind not in {"created", "delivered", "paid_date"}:
        raise ValueError("Неизвестный фильтр даты.")
    mode = args.get("mode", "pending")
    if mode not in set(LABELS) | {"all", "registries"}:
        mode = "pending"
    query = str(args.get("q") or "").strip()[:255]
    cohort = []
    for original in snapshot.get("rows", []):
        row = dict(original)
        if kind == "paid_date":
            entries = [e for e in row["entries"] if start <= e["date"] <= end]
            if not entries:
                continue
            row["net"] = str(sum((amount(e["net"]) for e in entries), Decimal(0)))
            row["entries"] = entries
        elif not row.get(kind) or not start <= row[kind] <= end:
            continue
        cohort.append(row)
    counts = {k: sum(r["category"] == k for r in cohort) for k in LABELS}
    counts["all"] = len(cohort)
    pending = [r for r in cohort if r["category"] == "pending"]
    expected = [r for r in cohort if r["category"] == "expected"]
    paid = [r for r in cohort if r["category"] == "paid"]
    def total(rows, key):
        return sum((amount(r[key]) for r in rows), Decimal(0))
    totals = dict(pending=total(pending, "remaining"), expected=total(expected, "cod"), paid=total(paid, "net"),
                  estimate=total(pending, "estimate") if all(r["estimate"] is not None for r in pending) else None)
    regs = [r for r in snapshot.get("registries", []) if start <= r["date"] <= end]
    counts["registries"] = len(regs)
    rows = ([r for r in regs if not query or query in r["number"]] if mode == "registries" else
            [r for r in cohort if (mode == "all" or r["category"] == mode)
             and (not query or query.casefold() in (r["track"]+" "+r["order"]).casefold())])
    rows.sort(key=lambda r: (r.get("date") or r.get("created", ""), r.get("track", "")), reverse=True)
    try:
        page = max(1, int(args.get("page", 1)))
    except (TypeError, ValueError):
        page = 1
    pages = max(1, (len(rows)+24)//25)
    page = min(page, pages)
    long_wait = [r for r in expected if (today-datetime.strptime(r["created"], "%Y-%m-%d").date()).days >= 30]
    return dict(start=start, end=end, date_kind=kind, mode=mode, query=query, rows=rows[(page-1)*25:page*25],
                page=page, pages=pages, total=len(rows), totals=totals, counts=counts,
                long_wait=long_wait, long_wait_sum=total(long_wait, "cod"),
                long_wait_percent=float(total(long_wait, "cod") / totals["expected"] * 100) if totals["expected"] else 0,
                coverage_gap=bool(snapshot and start < snapshot.get("since", start)))


def register_cdek_payouts_routes(app, service, load_references, allowed, csrf):
    sync = CdekSync(service)

    def authorize():
        if not allowed():
            abort(403)

    @app.get("/app/analytics/cdek-payouts")
    def cdek_payouts_page():
        authorize()
        error = ""
        try:
            snapshot, status = service.snapshot(), sync.summary()
        except CdekError as e:
            snapshot, status, error = {}, {}, str(e)
        try:
            model = build_view(snapshot, request.args, service.today())
        except ValueError as e:
            model, error = build_view(snapshot, {}, service.today()), str(e)
        def link(**changes):
            params = dict(model)
            params = dict(mode=model["mode"], date_kind=model["date_kind"], q=model["query"],
                          **{"from": model["start"], "to": model["end"]})
            params.update(changes)
            return url_for("cdek_payouts_page", **params)
        return render_template("cdek_payouts.html", payout=model, snapshot=snapshot, status=status,
                               error=error, money=money, link=link, labels=LABELS,
                               today=service.today().isoformat(), configured=service.client.configured,
                               checked=display_time(snapshot.get("checked_at", 0)),
                               stale=bool(snapshot and service.clock()-snapshot.get("checked_at", 0)>86400),
                               schedule=os_schedule())

    @app.route("/app/analytics/cdek-payouts/sync", methods=["GET", "POST"])
    def cdek_payouts_sync():
        authorize()
        try:
            if request.method == "POST":
                csrf()
                sync.start(load_references, app)
            return jsonify(sync.summary()), 202 if request.method == "POST" else 200
        except CdekError as e:
            return jsonify(message=str(e)), 409 if e.code == "CDEK_BUSY" else 503

    @app.post("/app/analytics/cdek-payouts/import")
    def cdek_payouts_import():
        authorize()
        csrf()
        upload = request.files.get("report")
        if not upload:
            return jsonify(message="Выберите отчёт XLSX."), 400
        try:
            tracks, since = import_report(upload.stream.read(10*1024*1024+1))
            if since > service.today().isoformat() or (service.today()-datetime.strptime(since, "%Y-%m-%d").date()).days > 730:
                raise ValueError("Допустимы накладные за последние два года.")
            lock = sync.lock()
            if not lock.acquire():
                raise CdekError("CDEK_BUSY", "Дождитесь завершения текущего обновления.")
            try:
                old = service.read("sources.json", {"tracks": [], "since": ""})
                service.write("sources.json", dict(tracks=sorted(set(old["tracks"]) | set(tracks)),
                              since=min(v for v in [since, old.get("since")] if v)))
                service.write("stage.json", {})
            finally:
                lock.release()
            return redirect(url_for("cdek_payouts_page"))
        except (ValueError, CdekError) as e:
            return jsonify(message=str(e)), 400


def os_schedule():
    import os
    return os.getenv("CDEK_PAYOUTS_TIMER_ENABLED") == "1"
