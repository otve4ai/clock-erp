"""Configurable, read-only XLSX export for the canonical product catalogue."""

import json
from collections import OrderedDict
from datetime import timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO

from app.catalog_db import CatalogDatabase
from app.time_ranking import parse_erp_datetime


FORMULA_PREFIXES = ("=", "+", "-", "@")

# Only fields backed by the current catalogue schema are offered. Warehouse
# columns are discovered separately from synchronized Bitrix payloads.
FIELD_GROUPS = (
    ("Основная информация", (
        ("sku", "Артикул / SKU"),
        ("name", "Название"),
        ("brand", "Бренд"),
        ("category", "Категория"),
        ("barcode", "Штрихкод"),
        ("description", "Описание"),
        ("status", "Статус"),
    )),
    ("Цены", (("price", "Цена"),)),
    ("Остатки", (("stock_total", "Общий остаток"),)),
    ("Дополнительно", (
        ("model", "Модель"),
        ("cell", "Ячейка"),
        ("created_at", "Дата создания"),
        ("updated_at", "Дата обновления"),
    )),
)
FIELD_LABELS = OrderedDict(
    field for _group, fields in FIELD_GROUPS for field in fields
)
DEFAULT_FIELDS = (
    "sku", "name", "brand", "category", "barcode", "price", "stock_total"
)


def safe_excel_text(value):
    text = str(value or "")
    return "'" + text if text.startswith(FORMULA_PREFIXES) else text


def excel_number(value):
    if value in (None, ""):
        return None
    try:
        return float(Decimal(str(value)))
    except (InvalidOperation, TypeError, ValueError):
        return None


def excel_datetime(value):
    parsed = parse_erp_datetime(value)
    if parsed is None:
        return None
    result = parsed[0]
    if parsed[1] == "date":
        return result.date()
    if result.tzinfo is not None:
        result = result.astimezone(timezone.utc).replace(tzinfo=None)
    return result


def _warehouse_rows(payload):
    """Normalize supported Bitrix warehouse payload shapes without fixed names."""
    try:
        source = json.loads(payload) if isinstance(payload, str) else payload
    except (TypeError, ValueError):
        return []
    source = source if isinstance(source, dict) else {}
    stocks = source.get("warehouse_stocks") or source.get("STORES") or []
    if isinstance(stocks, dict):
        rows = []
        for key, value in stocks.items():
            if isinstance(value, dict):
                row = dict(value)
                row.setdefault("name", key)
            else:
                row = {"name": key, "quantity": value}
            rows.append(row)
        return rows
    return [row for row in stocks if isinstance(row, dict)]


def warehouse_values(payload):
    result = {}
    for row in _warehouse_rows(payload):
        name = next((
            str(row.get(key) or "").strip()
            for key in ("name", "warehouse_name", "store_name", "title", "NAME")
            if str(row.get(key) or "").strip()
        ), "")
        if not name:
            continue
        raw_quantity = next((
            row.get(key) for key in
            ("quantity", "stock", "amount", "available", "value", "QUANTITY")
            if row.get(key) not in (None, "")
        ), 0)
        quantity = excel_number(raw_quantity)
        result[name] = quantity if quantity is not None else 0
    return result


class ProductExcelExport:
    def __init__(self, database=None):
        self.database = database or CatalogDatabase(cache_initialization=True)

    @staticmethod
    def field_groups():
        return [
            {"label": label, "fields": [
                {"key": key, "label": field_label, "default": key in DEFAULT_FIELDS}
                for key, field_label in fields
            ]}
            for label, fields in FIELD_GROUPS
        ]

    def available_warehouses(self):
        names = set()
        with self.database.connect() as connection:
            rows = connection.execute(
                "SELECT cp.normalized_payload_json FROM catalog_excel_products p "
                "JOIN catalog_excel_batches b ON b.id=p.current_batch_id "
                "JOIN catalog_products cp ON cp.id=p.bitrix_catalog_product_id "
                "WHERE p.active=1 AND (p.source_key LIKE 'bitrix:%' "
                "OR (b.status='active' AND p.current_batch_id=b.id))"
            ).fetchall()
        for row in rows:
            names.update(warehouse_values(row["normalized_payload_json"]))
        return sorted(names, key=str.casefold)

    def validate_fields(self, requested_fields, available_warehouses):
        requested = list(dict.fromkeys(str(value) for value in requested_fields or []))
        if not requested:
            requested = list(DEFAULT_FIELDS) + [
                "warehouse:" + name for name in available_warehouses
            ]
        allowed_warehouses = set(available_warehouses)
        validated = []
        for field in requested:
            if field in FIELD_LABELS:
                validated.append(field)
            elif field.startswith("warehouse:") and field[10:] in allowed_warehouses:
                validated.append(field)
            else:
                raise ValueError("Выбрано неизвестное поле экспорта.")
        if not validated:
            raise ValueError("Выберите хотя бы одно поле для экспорта.")
        return validated

    def enrich(self, products):
        """Attach synchronized per-warehouse balances with bounded queries."""
        products = list(products)
        ids = [int(product["id"]) for product in products]
        if not ids:
            return products
        payloads = {}
        with self.database.connect() as connection:
            for start in range(0, len(ids), 400):
                chunk = ids[start:start + 400]
                placeholders = ", ".join("?" for _ in chunk)
                rows = connection.execute(
                    "SELECT p.id AS product_id,cp.normalized_payload_json "
                    "FROM catalog_excel_products p LEFT JOIN catalog_products cp "
                    "ON cp.id=p.bitrix_catalog_product_id "
                    "WHERE p.id IN ({})".format(placeholders), chunk,
                ).fetchall()
                payloads.update({
                    int(row["product_id"]): row["normalized_payload_json"]
                    for row in rows if row["normalized_payload_json"]
                })
        for product in products:
            product["_export_warehouses"] = warehouse_values(
                payloads.get(int(product["id"]))
            )
        return products

    def build(self, products, total, fields=None, available_warehouses=None):
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        warehouses = list(available_warehouses or self.available_warehouses())
        fields = self.validate_fields(fields, warehouses)
        headers = [
            FIELD_LABELS.get(field, field[10:] if field.startswith("warehouse:") else field)
            for field in fields
        ]

        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("Товары")
        sheet.freeze_panes = "A2"
        last_column = get_column_letter(len(fields))
        sheet.auto_filter.ref = "A1:{}{}".format(last_column, max(1, int(total) + 1))
        header = []
        for index, (field, header_text) in enumerate(zip(fields, headers), 1):
            width = 18
            if field in {"name", "description"}:
                width = 44
            elif field in {"brand", "category"} or field.startswith("warehouse:"):
                width = 22
            elif field in {"created_at", "updated_at"}:
                width = 21
            sheet.column_dimensions[get_column_letter(index)].width = width
            cell = WriteOnlyCell(sheet, value=header_text)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="174887")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            header.append(cell)
        sheet.append(header)

        for product in products:
            row = []
            for field in fields:
                value = self._value(product, field)
                cell = WriteOnlyCell(sheet, value=value)
                if field in {"price", "stock_total"} or field.startswith("warehouse:"):
                    cell.number_format = '#,##0.00' if field == "price" else '#,##0.###'
                elif field in {"created_at", "updated_at"} and value is not None:
                    cell.number_format = "DD.MM.YYYY HH:MM"
                cell.alignment = Alignment(vertical="top", wrap_text=field == "description")
                row.append(cell)
            sheet.append(row)

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def _value(product, field):
        if field.startswith("warehouse:"):
            return excel_number(
                (product.get("_export_warehouses") or {}).get(field[10:], 0)
            ) or 0
        values = {
            "sku": safe_excel_text(product.get("excel_article")),
            "name": safe_excel_text(product.get("display_name") or product.get("excel_name_raw")),
            "brand": safe_excel_text(product.get("display_brand") or product.get("excel_brand")),
            "category": safe_excel_text(product.get("display_category") or product.get("excel_category")),
            "barcode": safe_excel_text(product.get("bitrix_barcode")),
            "description": safe_excel_text(product.get("bitrix_description")),
            "status": "Активен" if product.get("active") else "Неактивен",
            "price": excel_number(product.get("bitrix_price_amount")),
            "stock_total": excel_number(product.get("stock")) or 0,
            "model": safe_excel_text(product.get("model")),
            "cell": safe_excel_text(product.get("cell")),
            "created_at": excel_datetime(product.get("created_at")),
            "updated_at": excel_datetime(product.get("updated_at")),
        }
        return values[field]
